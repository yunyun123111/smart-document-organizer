"""Phase 3 文件解析层测试：PDF/图片/DOCX/XLSX/TXT/CSV + 文本预处理。

测试前自动在 tmp 目录生成各类型样本文件。
"""
from __future__ import annotations

from pathlib import Path

import pymupdf as fitz
import pytest
from openpyxl import Workbook
from PIL import Image, ImageDraw
from docx import Document as DocxDocument

from backend.services.parser_service import ParserService, ParsedDocument
from backend.services.text_service import clean_text, clean_parsed_document


# ============ 样本生成 ============
@pytest.fixture
def sample_dir(tmp_path: Path) -> Path:
    return tmp_path / "samples"


def make_text_pdf(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page()
    lines = [
        "销售合同",
        "合同编号：XS202608001",
        "甲方：ABC有限公司",
        "乙方：DEF有限公司",
        "合同金额：人民币 128,500.00 元",
        "签订日期：2026年08月20日",
        "本合同由甲乙双方根据《中华人民共和国民法典》订立，双方同意按照本合同条款执行。",
        "付款方式：合同签订后支付50%，验收合格后支付剩余50%。",
        "交货地点：甲方指定仓库。",
    ]
    y = 72
    for line in lines:
        # fontname="china-s" 使用内置简体中文字体，避免中文变占位符
        page.insert_text((72, y), line, fontname="china-s")
        y += 24
    doc.save(str(path))
    doc.close()


def make_scanned_pdf(path: Path) -> None:
    """无文本层的"扫描件"：插入一张图片，无文本。"""
    doc = fitz.open()
    page = doc.new_page()
    img = Image.new("RGB", (200, 100), (255, 255, 255))
    img.save(str(path.with_suffix(".png")))
    page.insert_image(page.rect, filename=str(path.with_suffix(".png")))
    doc.save(str(path))
    doc.close()
    path.with_suffix(".png").unlink()


def make_image(path: Path) -> None:
    img = Image.new("RGB", (300, 150), "white")
    draw = ImageDraw.Draw(img)
    draw.text((20, 40), "增值税发票 发票号码 12345678", fill="black")
    draw.text((20, 80), "金额 128,500.00", fill="black")
    img.save(str(path))


def make_docx(path: Path) -> None:
    doc = DocxDocument()
    doc.add_heading("采购合同", level=1)
    doc.add_paragraph("甲方：ABC有限公司")
    doc.add_paragraph("乙方：DEF有限公司")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "采购编号"
    table.cell(0, 1).text = "CG202608001"
    table.cell(1, 0).text = "金额"
    table.cell(1, 1).text = "50000"
    doc.save(str(path))


def make_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "销售订单"
    ws.append(["客户名称", "订单编号", "日期", "金额"])
    ws.append(["ABC公司", "DD001", "2026-08-20", 128500])
    ws2 = wb.create_sheet("收款记录")
    ws2.append(["收款日期", "金额"])
    ws2.append(["2026-08-21", 20000])
    wb.save(str(path))


@pytest.fixture
def samples(sample_dir: Path):
    sample_dir.mkdir(parents=True, exist_ok=True)
    files: dict[str, Path] = {}
    p = sample_dir / "text.pdf"
    make_text_pdf(p)
    files["text_pdf"] = p
    p = sample_dir / "scan.pdf"
    make_scanned_pdf(p)
    files["scanned_pdf"] = p
    p = sample_dir / "img.jpg"
    make_image(p)
    files["jpg"] = p
    p = sample_dir / "img.png"
    make_image(p)
    files["png"] = p
    p = sample_dir / "contract.docx"
    make_docx(p)
    files["docx"] = p
    p = sample_dir / "orders.xlsx"
    make_xlsx(p)
    files["xlsx"] = p
    p = sample_dir / "notes.txt"
    p.write_text("2026 年 08 月 20 日 对账说明\n金额：人民币 128,500.00 元", encoding="utf-8")
    files["txt"] = p
    p = sample_dir / "data.csv"
    p.write_text("客户名称,订单编号,日期,金额\nABC公司,DD001,2026-08-20,128500", encoding="utf-8")
    files["csv"] = p
    return files


# ============ 测试 ============
class TestPdfParser:
    def test_text_pdf(self, samples):
        result = parser.parse_file(samples["text_pdf"])
        assert isinstance(result, ParsedDocument)
        assert result.needs_ocr is False
        assert result.pages == 1
        assert "销售合同" in result.text
        assert result.metadata["page_count"] == 1

    def test_scanned_pdf_flagged_for_ocr(self, samples):
        result = parser.parse_file(samples["scanned_pdf"])
        assert result.needs_ocr is True
        assert result.text == ""
        assert result.pages == 1

    def test_corrupted_pdf(self, sample_dir):
        bad = sample_dir.parent / "bad.pdf"
        bad.write_bytes(b"%PDF-1.4 not a real pdf")
        with pytest.raises(Exception):
            parser.parse_file(bad)


class TestImageParser:
    @pytest.mark.parametrize("key", ["jpg", "png"])
    def test_image_flagged_for_ocr(self, samples, key):
        result = parser.parse_file(samples[key])
        assert result.needs_ocr is True
        assert result.metadata["width"] == 300
        assert result.metadata["height"] == 150
        assert len(result.images) == 1


class TestWordParser:
    def test_docx_content(self, samples):
        result = parser.parse_file(samples["docx"])
        assert "[标题] 采购合同" in result.text
        assert "甲方：ABC有限公司" in result.text
        assert "采购编号 | CG202608001" in result.text
        assert len(result.tables) == 1
        assert result.tables[0][0] == ["采购编号", "CG202608001"]


class TestExcelParser:
    def test_xlsx_content(self, samples):
        result = parser.parse_file(samples["xlsx"])
        assert "Sheet: 销售订单" in result.text
        assert "客户名称 | 订单编号 | 日期 | 金额" in result.text
        assert "ABC公司 | DD001 | 2026-08-20 | 128500" in result.text
        assert len(result.metadata["sheets"]) == 2
        assert len(result.tables) == 2


class TestTextParsers:
    def test_txt(self, samples):
        result = parser.parse_file(samples["txt"])
        assert "对账说明" in result.text
        assert result.metadata["lines"] == 2

    def test_csv(self, samples):
        result = parser.parse_file(samples["csv"])
        assert len(result.tables) == 1
        assert result.tables[0][0] == ["客户名称", "订单编号", "日期", "金额"]
        assert result.tables[0][1][1] == "DD001"

    def test_csv_semicolon_delimiter(self, tmp_path):
        p = tmp_path / "semi.csv"
        p.write_text("名称;数量;单价\nA;2;10", encoding="utf-8")
        result = parser.parse_file(p)
        assert result.tables[0][1] == ["A", "2", "10"]


class TestParserService:
    def test_unsupported_type(self, tmp_path):
        p = tmp_path / "evil.exe"
        p.write_bytes(b"MZ")
        with pytest.raises(ValueError, match="不支持"):
            parser.parse_file(p)

    def test_missing_file(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            parser.parse_file(tmp_path / "nope.pdf")

    def test_case_insensitive_extension(self, samples):
        # 大写扩展名也应识别
        p = samples["txt"].with_suffix(".TXT")
        p.write_text("大写扩展名测试", encoding="utf-8")
        result = parser.parse_file(p)
        assert "大写扩展名测试" in result.text


class TestTextService:
    def test_normalize_date(self):
        assert clean_text("2026 年 08 月 20 日") == "2026-08-20"
        assert clean_text("合同日期 2026/8/5 签订") == "合同日期 2026-08-05 签订"

    def test_normalize_amount(self):
        assert clean_text("人民币 128,500.00 元") == "128500.00"
        assert clean_text("金额 ¥128500") == "金额 128500"

    def test_whitespace_and_newlines(self):
        assert clean_text("甲  方   \t 乙\n\n\n\n丙") == "甲 方 乙\n\n丙"

    def test_duplicate_chars(self):
        assert clean_text("！！！！！重要") == "！！！重要"

    def test_clean_parsed_document_tables(self, samples):
        result = parser.parse_file(samples["csv"])
        cleaned = clean_parsed_document(result)
        assert cleaned.tables[0][1][3] == "128500"


parser = ParserService()
