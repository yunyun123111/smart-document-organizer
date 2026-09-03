"""Excel 登记表权威数据源匹配服务。

把用户手工登记的 Excel 台账（合同 / 结算 / 收付款等）作为权威主数据：
OCR 提取字段后，用合同号（主键，支持后缀容错）或兜底键（日期+船名 / 物料+数量）
匹配 Excel 行，命中则用 Excel 行字段覆盖 OCR 结果（置信度 1.0，来源 EXCEL），
再走命名模板归档。未命中完全回退现有流程，不消耗额外 AI。

设计要点：
- 一个工作簿（excel_sources）可含多个 sheet（excel_sheet_configs），每个 sheet 独立启用 / 列映射
- 列映射自动识别（标准字段 -> Excel 列名），支持用户手动修正
- 匹配全在本地内存索引完成，零 token；索引按文件内容哈希 + 配置变更自动重建
"""
from __future__ import annotations

import json
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from sqlalchemy.orm import Session

from backend.models.excel_source import ExcelSheetConfig, ExcelSource
from backend.utils.logger import get_logger

logger = get_logger("services.excel_matcher")

# 与 field_extractor.FIELD_ORDER 对齐的标准字段集合（EXCEL 可覆盖的字段）
STANDARD_FIELDS = [
    "company", "seller", "date", "contract_no", "contract_suffix",
    "order_no", "invoice_no", "amount", "vessel", "material", "quantity", "unit_price",
]

# 标准字段 -> 可能的 Excel 列名（顺序即优先级，靠前者优先）
# 列名先经 _norm_col 规范化（去空白/标点/括号、转小写）再比对
COLUMN_ALIASES: dict[str, list[str]] = {
    "contract_no": [
        "销售合同编号oa", "销售合同编号", "采购合同编号oa", "采购合同编号",
        "合同编号", "合同号", "销售合同号", "采购合同号",
    ],
    "vessel": ["船名", "船号"],
    "material": ["物料名称", "物料", "品名", "商品名称"],
    "quantity": ["数量", "结算数量", "销售量", "采购量", "已收货转量", "采购结算数量", "销售结算数"],
    "unit_price": ["单价", "采购单价", "销售单价"],
    "amount": [
        "金额", "采购合同金额", "销售合同金额", "采购结算金额", "销售结算金额",
        "合同金额", "结算金额", "回款金额", "支付金额", "销售金额",
    ],
    "company": ["供应商", "供货方", "卖方"],
    "seller": ["采购商名称", "销售客户", "采购商", "客户名称"],
    "date": ["日期", "采购合同签订日期", "销售合同签订日期", "开票日期", "签约日期"],
    "order_no": ["2.0编号", "订单号", "外部编号"],
    "invoice_no": ["发票号", "发票号码", "发票代码"],
}

# 匹配键：Excel 中合同号列名（自动识别时主键列优先找合同号）
CONTRACT_KEY_COLUMNS = ["销售合同编号oa", "销售合同编号", "采购合同编号oa", "采购合同编号", "合同编号", "合同号"]


def _norm_col(name: str) -> str:
    """规范化列名用于匹配：去空白/标点/括号、统一小写。"""
    if not name:
        return ""
    s = str(name)
    s = unicodedata.normalize("NFKC", s)
    # 全角括号/空格转半角
    s = s.replace("（", "(").replace("）", ")").replace("　", "")
    # 去空白与常见标点
    out = []
    for ch in s:
        if ch in " \t\u3000":
            continue
        if ch in "()（）[]【】-_/\\.:：,，;；、":
            continue
        out.append(ch.lower())
    return "".join(out)


def _norm_key(s: str) -> str:
    """规范化匹配键（合同号 / 船名 / 物料）：去空白、全角转半角、统一括号、转大写。"""
    if s is None:
        return ""
    s = str(s)
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("（", "(").replace("）", ")").replace("　", " ").strip()
    s = "".join(ch for ch in s if not ch.isspace())
    return s.upper()


def _norm_qty(v) -> str:
    """数量/金额规范化：去千分位逗号，统一为纯数字串。"""
    if v is None:
        return ""
    s = str(v).replace(",", "").replace("，", "").strip()
    return s


def _norm_date(v) -> str:
    """日期规范化：datetime/date -> YYYY-MM-DD；字符串尝试解析。"""
    if v is None or v == "":
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    # 去掉时间部分
    s = s.split(" ")[0].split("T")[0]
    # 兼容 2026/2/6 2026.2.6 2026年2月6日
    for sep in ("/", ".", "-"):
        if sep in s:
            parts = s.split(sep)
            if len(parts) == 3 and all(x.isdigit() for x in parts):
                return f"{int(parts[0]):04d}-{int(parts[1]):02d}-{int(parts[2]):02d}"
    if "年" in s and "月" in s:
        import re
        m = re.match(r"(\d{4})年(\d{1,2})月(\d{1,2})?日?", s)
        if m:
            y, mo = int(m.group(1)), int(m.group(2))
            d = int(m.group(3)) if m.group(3) else 1
            return f"{y:04d}-{mo:02d}-{d:02d}"
    return s


def detect_column_map(headers: list[str]) -> dict[str, str]:
    """根据表头自动识别标准字段 -> Excel 列名。返回 {标准字段: 原列名}。

    同一台账常同时含采购侧与销售侧（采购合同金额 / 销售合同金额、两个单价列），
    额外识别 *_sale 侧字段，供按合同前缀（SJWLXS/HSYCXS=销售）选择对应侧。
    """
    normalized = {_norm_col(h): h for h in headers if h}
    result: dict[str, str] = {}
    for std_field, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                result[std_field] = normalized[alias]
                break

    # 销售侧专用字段（采购侧为主，销售侧单独存 *_sale）
    for alias in ("销售合同金额", "销售结算金额", "销售金额"):
        key = _norm_col(alias)
        if key in normalized:
            result["amount_sale"] = normalized[key]
            break
    for alias in ("销售合同签订日期", "销售日期"):
        key = _norm_col(alias)
        if key in normalized:
            result["date_sale"] = normalized[key]
            break
    for alias in ("销售单价",):
        key = _norm_col(alias)
        if key in normalized:
            result["unit_price_sale"] = normalized[key]
            break
    # 表头出现两个"单价"列：第二个视为销售单价
    if "unit_price_sale" not in result:
        if [(_norm_col(h)) for h in headers if h].count("单价") >= 2:
            seen = 0
            for h in headers:
                if h is not None and _norm_col(h) == "单价":
                    seen += 1
                    if seen == 2:
                        result["unit_price_sale"] = h
                        break
    return result


def detect_key_column(headers: list[str]) -> str:
    """自动识别合同号主键列名（返回原列名；无则空串）。"""
    normalized = {_norm_col(h): h for h in headers if h}
    for alias in CONTRACT_KEY_COLUMNS:
        if alias in normalized:
            return normalized[alias]
    return ""


@dataclass
class ExcelMatchResult:
    """一次 Excel 匹配的结果。"""
    matched: bool = False
    source_name: str = ""
    sheet_name: str = ""
    key: str = ""                    # 匹配键说明（用于人工审核展示）
    matched_field: str = ""          # contract_no / vessel_date / material_quantity
    row_index: int = -1
    fields: dict = field(default_factory=dict)   # 覆盖字段：标准字段 -> 值

    def to_dict(self) -> dict:
        return {
            "matched": self.matched,
            "source_name": self.source_name,
            "sheet_name": self.sheet_name,
            "key": self.key,
            "matched_field": self.matched_field,
            "fields": self.fields,
        }


class _LoadedSheet:
    """一个已加载并建好索引的 sheet。"""
    def __init__(self, cfg: ExcelSheetConfig, rows: list[dict], source_name: str = ""):
        self.cfg = cfg
        self.source_name = source_name
        self.rows = rows
        # 主索引：规范合同号 -> [row_idx]
        self.by_contract: dict[str, list[int]] = {}
        # 兜底索引：vessel_date (norm_date, norm_vessel) -> [row_idx]
        self.by_vessel_date: dict[tuple[str, str], list[int]] = {}
        # 兜底索引：material_quantity (norm_material, norm_qty) -> [row_idx]
        self.by_material_qty: dict[tuple[str, str], list[int]] = {}
        self._build()

    def _build(self) -> None:
        for i, row in enumerate(self.rows):
            # 主键列合同号 + 备选合同号（销售/采购等所有合同编号列）
            cns = set()
            primary = _norm_key(row.get("contract_no", ""))
            if primary:
                cns.add(primary)
            for extra in row.get("__contracts", []) or []:
                e = _norm_key(extra)
                if e:
                    cns.add(e)
            for cn in cns:
                self.by_contract.setdefault(cn, []).append(i)
            d = _norm_date(row.get("date", ""))
            v = _norm_key(row.get("vessel", ""))
            if d and v:
                self.by_vessel_date.setdefault((d, v), []).append(i)
            m = _norm_key(row.get("material", ""))
            q = _norm_qty(row.get("quantity", ""))
            if m and q:
                self.by_material_qty.setdefault((m, q), []).append(i)


class ExcelMatcher:
    """Excel 数据源管理器：加载、索引缓存、匹配。"""

    def __init__(self) -> None:
        # source_id -> list[_LoadedSheet]（仅启用 sheet）
        self._cache: dict[int, list[_LoadedSheet]] = {}
        self._cache_sig: dict[int, str] = {}   # source_id -> 配置签名（列映射+行数+hash）

    # ---------- 缓存管理 ----------
    def _config_signature(self, db: Session, src: ExcelSource) -> str:
        sig = [src.file_hash or str(Path(src.file_path).stat().st_mtime_ns) if Path(src.file_path).exists() else ""]
        for s in src.sheets:
            sig.append(f"{s.id}:{s.enabled}:{s.key_column}:{s.column_map}:{s.row_count}")
        return "|".join(sig)

    def reload_source(self, db: Session, source_id: int) -> None:
        """强制重建某个数据源的索引（上传/配置变更后调用）。"""
        src = db.get(ExcelSource, source_id)
        if src is None:
            self._cache.pop(source_id, None)
            self._cache_sig.pop(source_id, None)
            return
        loaded: list[_LoadedSheet] = []
        for cfg in src.sheets:
            if not cfg.enabled:
                continue
            try:
                rows = self._parse_sheet(Path(src.file_path), cfg)
                loaded.append(_LoadedSheet(cfg, rows, src.name))
                cfg.last_error = ""
            except Exception as e:  # noqa: BLE001
                logger.warning("解析 Excel sheet %s/%s 失败: %s", src.name, cfg.sheet_name, e)
                cfg.last_error = str(e)[:480]
        self._cache[source_id] = loaded
        self._cache_sig[source_id] = self._config_signature(db, src)
        db.commit()
        logger.info("Excel 数据源 #%s 索引重建完成（%d 个启用 sheet）", source_id, len(loaded))

    def _ensure_loaded(self, db: Session, src: ExcelSource) -> list[_LoadedSheet]:
        sig = self._config_signature(db, src)
        if src.id in self._cache and self._cache_sig.get(src.id) == sig:
            return self._cache[src.id]
        self.reload_source(db, src.id)
        return self._cache.get(src.id, [])

    def _parse_sheet(self, path: Path, cfg: ExcelSheetConfig) -> list[dict]:
        """按配置的列映射把 sheet 解析为标准字段行列表。"""
        import openpyxl
        try:
            col_map = json.loads(cfg.column_map) if cfg.column_map else {}
        except Exception:  # noqa: BLE001
            col_map = {}
        if not col_map:
            return []

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            if cfg.sheet_name not in wb.sheetnames:
                raise ValueError(f"sheet 不存在: {cfg.sheet_name}")
            ws = wb[cfg.sheet_name]
            all_rows = list(ws.iter_rows(values_only=True))
            if not all_rows:
                return []

            # 表头定位：前 10 行内找能命中 >=2 个映射列的一行
            header_names: list = []
            header_map: dict[str, list[int]] = {}
            header_idx = -1
            for i, r in enumerate(all_rows[:10]):
                norm = {_norm_col(str(c)) if c is not None else "": c for c in r}
                hits = sum(1 for v in col_map.values() if v and _norm_col(str(v)) in norm)
                if hits >= 2:
                    header_names = list(r)
                    for i2, c in enumerate(r):
                        if c is not None:
                            header_map.setdefault(_norm_col(str(c)), []).append(i2)
                    header_idx = i
                    break
            if header_idx < 0 or not header_map:
                return []

            # 标准字段 -> 列下标（同名列保留多位置：unit_price_sale 取第二个"单价"）
            std_col_idx: dict[str, int] = {}
            for std, excel_col in col_map.items():
                if not excel_col:
                    continue
                idxs = header_map.get(_norm_col(str(excel_col)), [])
                if not idxs:
                    continue
                idx = idxs[0]
                if std == "unit_price_sale" and col_map.get("unit_price") == excel_col and len(idxs) > 1:
                    idx = idxs[1]
                std_col_idx[std] = idx
            if not std_col_idx:
                return []

            # 收集所有"合同编号/合同号"列（含主键列），作为备选匹配键
            contract_cols: list[int] = []
            for h, idxs in header_map.items():
                if "合同" in h and ("编号" in h or h.endswith("合同号")):
                    contract_cols.extend(idxs)

            rows: list[dict] = []
            for row in all_rows[header_idx + 1:]:
                vals: dict = {}
                for std, idx in std_col_idx.items():
                    if idx < len(row):
                        v = row[idx]
                        if v is not None and str(v).strip() != "":
                            vals[std] = _fmt_value(std, v)
                contracts: list[str] = []
                for ci in contract_cols:
                    if ci < len(row) and row[ci] is not None and str(row[ci]).strip() != "":
                        contracts.append(_fmt_value("contract_no", row[ci]))
                if contracts:
                    vals["__contracts"] = contracts
                if vals:
                    rows.append(vals)
            return rows
        finally:
            wb.close()

    # ---------- 匹配 ----------
    def match(self, db: Session, field_values: dict) -> ExcelMatchResult:
        """根据识别出的字段（field_values）在所有启用数据源中匹配。"""
        srcs = db.query(ExcelSource).filter(ExcelSource.enabled.is_(True)).all()
        if not srcs:
            return ExcelMatchResult()

        ocr_cn = _norm_key(field_values.get("contract_no", ""))
        ocr_date = _norm_date(field_values.get("date", ""))
        ocr_vessel = _norm_key(field_values.get("vessel", ""))
        ocr_material = _norm_key(field_values.get("material", ""))
        ocr_qty = _norm_qty(field_values.get("quantity", ""))

        for src in srcs:
            sheets = self._ensure_loaded(db, src)
            for sheet in sheets:
                # 1. 主键：合同号精确匹配
                if ocr_cn:
                    r = self._match_contract(sheet, ocr_cn, field_values)
                    if r.matched:
                        return r
                # 2. 主键后缀容错（OCR 只识别到编号末段）
                if ocr_cn and len(_alnum(ocr_cn)) >= 3:
                    r = self._match_contract_suffix(sheet, ocr_cn, field_values)
                    if r.matched:
                        return r
                # 3. 兜底：日期 + 船名
                if ocr_date and ocr_vessel:
                    r = self._match_vessel_date(sheet, ocr_date, ocr_vessel, field_values)
                    if r.matched:
                        return r
                # 4. 兜底：物料 + 数量
                if ocr_material and ocr_qty:
                    r = self._match_material_qty(sheet, ocr_material, ocr_qty, field_values)
                    if r.matched:
                        return r
        return ExcelMatchResult()

    def _match_contract(self, sheet: _LoadedSheet, ocr_cn: str, fv: dict) -> ExcelMatchResult:
        idxs = sheet.by_contract.get(ocr_cn, [])
        if len(idxs) == 1:
            return self._build_result(sheet, idxs[0], "contract_no", f"合同号: {ocr_cn}", fv)
        if len(idxs) > 1:
            # 多行命中：取第一个并记录
            return self._build_result(sheet, idxs[0], "contract_no", f"合同号(多行): {ocr_cn}", fv)
        return ExcelMatchResult()

    def _match_contract_suffix(self, sheet: _LoadedSheet, ocr_cn: str, fv: dict) -> ExcelMatchResult:
        """后缀匹配：OCR 只识别到编号末段（如 YC0453）。取末 4-6 位字母数字。"""
        tail = _alnum(ocr_cn)
        if len(tail) < 3:
            return ExcelMatchResult()
        for size in (6, 5, 4):
            if len(tail) < size:
                continue
            suf = tail[-size:]
            hits = [i for i in range(len(sheet.rows)) if _norm_key(sheet.rows[i].get("contract_no", "")).endswith(suf)]
            if len(hits) == 1:
                return self._build_result(sheet, hits[0], "contract_no", f"合同号末段: {suf}", fv)
            if len(hits) > 1:
                # 多行命中后缀 -> 不冒险，回退
                return ExcelMatchResult()
        return ExcelMatchResult()

    def _match_vessel_date(self, sheet: _LoadedSheet, d: str, v: str, fv: dict) -> ExcelMatchResult:
        idxs = sheet.by_vessel_date.get((d, v), [])
        if len(idxs) == 1:
            return self._build_result(sheet, idxs[0], "vessel_date", f"日期+船名: {d} {v}", fv)
        return ExcelMatchResult()

    def _match_material_qty(self, sheet: _LoadedSheet, m: str, q: str, fv: dict) -> ExcelMatchResult:
        idxs = sheet.by_material_qty.get((m, q), [])
        if len(idxs) == 1:
            return self._build_result(sheet, idxs[0], "material_quantity", f"物料+数量: {m} {q}", fv)
        return ExcelMatchResult()

    def _build_result(self, sheet: _LoadedSheet, row_idx: int, matched_field: str, key: str, fv: dict) -> ExcelMatchResult:
        row = sheet.rows[row_idx]
        # 合同前缀判销售/采购侧：SJWLXS / HSYCXS = 销售合同，用销售侧金额/单价/日期
        _cn = str(row.get("contract_no", "") or "").upper()
        _is_sale = ("SJWLXS" in _cn) or ("HSYCXS" in _cn)
        # 只覆盖 Excel 行中存在的标准字段（与 OCR 字段合并，EXCEL 优先）
        overrides: dict[str, str] = {}
        for k, v in row.items():
            if k.startswith("_"):
                continue  # 跳过 __contracts 等内部字段
            if v is None or str(v).strip() == "":
                continue
            if k.endswith("_sale"):
                # 销售侧值：仅销售合同使用
                if _is_sale:
                    overrides[k[:-5]] = str(v)
                continue
            # 普通键（采购侧）：若为销售合同且存在销售侧值，让销售侧覆盖
            if _is_sale and k in ("amount", "unit_price", "date"):
                if f"{k}_sale" in row and str(row.get(f"{k}_sale", "") or "").strip() != "":
                    continue
            overrides[k] = str(v)
        return ExcelMatchResult(
            matched=True,
            source_name=sheet.source_name,
            sheet_name=sheet.cfg.sheet_name,
            key=key,
            matched_field=matched_field,
            row_index=row_idx,
            fields=overrides,
        )


def _alnum(s: str) -> str:
    """提取字母数字串。"""
    return "".join(ch for ch in s if ch.isalnum()).upper()


def _fmt_value(std_field: str, v) -> str:
    """单元格值 -> 标准字符串。"""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        # 整数金额/数量去掉 .0
        if v == int(v):
            return str(int(v))
        return str(v)
    return str(v).strip()


excel_matcher = ExcelMatcher()
