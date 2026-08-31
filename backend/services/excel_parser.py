"""Excel (XLSX) 解析器（规格书第二十一节）。

读取 Workbook → Sheet → 表头 → 数据，生成结构化文本交给规则/AI 分析。
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from backend.services.parser_service import BaseParser, ParsedDocument
from backend.utils.logger import get_logger

logger = get_logger("services.excel_parser")

# 每 sheet 最多提取的行数（防止超大表把内存和文本撑爆）
_MAX_ROWS_PER_SHEET = 500


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


class ExcelParser(BaseParser):
    file_types = ("xlsx",)

    def parse(self, file_path: str | Path) -> ParsedDocument:
        path = Path(file_path)
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            lines: list[str] = []
            tables: list[list[list[str]]] = []
            sheet_meta: dict[str, int] = {}

            for ws in wb.worksheets:
                sheet_meta[ws.title] = ws.max_row or 0
                lines.append(f"Sheet: {ws.title}")
                sheet_rows: list[list[str]] = []
                for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    if row_idx > _MAX_ROWS_PER_SHEET:
                        lines.append(f"…（已截断，该表超过 {_MAX_ROWS_PER_SHEET} 行）")
                        break
                    cells = [_cell_text(v) for v in row]
                    if any(cells):
                        sheet_rows.append(cells)
                tables.append(sheet_rows)
                if sheet_rows:
                    lines.append(" | ".join(sheet_rows[0]))  # 表头
                    for data_row in sheet_rows[1:]:
                        lines.append(" | ".join(data_row))
                lines.append("")

            text = "\n".join(lines).strip()
            parsed = ParsedDocument(
                text=text,
                metadata={
                    "sheets": list(sheet_meta.keys()),
                    "sheet_row_counts": sheet_meta,
                    "table_count": len(tables),
                },
                tables=tables,
            )
            logger.info("Excel 解析完成：%d 个 sheet，文本 %d 字符", len(sheet_meta), len(text))
            return parsed
        finally:
            wb.close()
